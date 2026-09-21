"""Level 4: Frontier broadband-plans automation using Patchright.

Drives system-installed Chrome for authentic TLS/JA4 fingerprint.
Combined with Bezier mouse, log-normal typing, and warm-up browsing.

Usage:
    python frontier_patchright.py [--proxy URL]
"""

from __future__ import annotations

import argparse
import os
import random
import time
from datetime import datetime

import openpyxl
from patchright.sync_api import (
    Browser,
    BrowserContext,
    Page,
    Playwright,
    TimeoutError as PatchrightTimeoutError,
    sync_playwright,
)

from api_logger import ApiLogger
from behavior import (
    bezier_click,
    lognormal_type,
    natural_scroll,
    reading_pause,
    warmup_browse,
)

FRONTIER_URL = "https://frontier.com/shop/internet"
ADDRESS = "2727 LBJ Freeway, Dallas, TX 75234"
LOGS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "logs", "frontier")
LEVEL_NAME = "frontier_level4_patchright"


def dismiss_cookie_banner(page: Page) -> None:
    for selector in [
        "button:has-text('Close')",
        "button:has-text('Accept All')",
        "button:has-text('Accept all')",
        "#truste-consent-button",
    ]:
        try:
            for loc in page.locator(selector).all():
                if loc.is_visible():
                    loc.click(timeout=2_000)
                    page.wait_for_timeout(500)
                    return
        except Exception:
            continue


def enter_address_landing_page(page: Page, api_logger: ApiLogger) -> None:
    print("  Entering address on landing page...")
    api_logger.set_phase("landing:address_entry")

    address_field = page.locator(
        "input[aria-label*='Street Address' i], "
        "input[placeholder*='Enter your address' i]"
    ).first
    address_field.wait_for(state="visible", timeout=15_000)
    bezier_click(page, address_field)
    address_field.fill("")

    lognormal_type(page, address_field, ADDRESS)
    reading_pause(1.5, 3.0)

    api_logger.set_phase("landing:autocomplete_select")
    suggestion = page.locator(
        "button:has-text('Dallas'), button:has-text('Johnson'), "
        "[role='option']:has-text('Dallas')"
    ).first
    try:
        suggestion.wait_for(state="visible", timeout=5_000)
        reading_pause(0.5, 1.5)
        bezier_click(page, suggestion)
        print("    selected autocomplete suggestion")
    except PatchrightTimeoutError:
        print("    no autocomplete — using fill fallback")
        address_field.fill(ADDRESS)

    reading_pause(0.5, 1.5)

    api_logger.set_phase("landing:check_availability_click")
    check_btn = page.locator("button:has-text('Check availability')").first
    try:
        check_btn.wait_for(state="visible", timeout=10_000)
        bezier_click(page, check_btn)
        print("    clicked 'Check availability'")
    except PatchrightTimeoutError:
        print("    Check availability button not found, pressing Enter")
        address_field.press("Enter")


def try_enter_address_landing(page: Page, api_logger: ApiLogger) -> bool:
    try:
        enter_address_landing_page(page, api_logger)
        return True
    except (PatchrightTimeoutError, Exception) as exc:
        print(f"    landing page address failed: {exc}")
        return False


def enter_address_buy_page(page: Page, api_logger: ApiLogger) -> None:
    print("  Entering address on buy page...")
    api_logger.set_phase("buy_page:address_entry")

    address_input = page.locator(
        "[role='combobox'], "
        "input[aria-label*='address' i], "
        "input[placeholder*='Enter your address' i]"
    ).first
    try:
        address_input.wait_for(state="visible", timeout=10_000)
    except PatchrightTimeoutError:
        print("    no address field on buy page — plans may already be shown")
        return

    bezier_click(page, address_input)
    address_input.fill("")

    lognormal_type(page, address_input, ADDRESS)
    reading_pause(2.0, 4.0)

    api_logger.set_phase("buy_page:autocomplete_select")
    option = page.locator("[role='option']").first
    try:
        option.wait_for(state="visible", timeout=8_000)
        reading_pause(0.5, 1.5)
        bezier_click(page, option)
        print("    selected suggestion on buy page")
    except PatchrightTimeoutError:
        print("    no suggestions appeared on buy page")
        address_input.fill(ADDRESS)

    reading_pause(0.5, 1.5)

    api_logger.set_phase("buy_page:check_availability_click")
    check_btn = page.locator("button:has-text('Check availability')").first
    try:
        check_btn.wait_for(state="visible", timeout=10_000)
        bezier_click(page, check_btn)
        print("    clicked 'Check availability' on buy page")
    except PatchrightTimeoutError:
        print("    no Check availability button")

    api_logger.set_phase("buy_page:plans_loading")
    page.wait_for_timeout(3_000)


def browse_plans(page: Page, api_logger: ApiLogger) -> None:
    api_logger.set_phase("plans:browsing")
    plan_selectors = [
        "button:has-text('View plan')",
        "button:has-text('Shop now')",
        "button:has-text('Add to cart')",
        "a:has-text('Fiber 500')",
        "a:has-text('Fiber 1 Gig')",
        "a:has-text('Fiber 2 Gig')",
        "[data-testid*='plan-card'] button",
        ".plan-card button",
    ]
    for selector in plan_selectors:
        try:
            buttons = page.locator(selector)
            count = buttons.count()
            if count > 0:
                print(f"  Found {count} plan button(s) matching '{selector}'")
                for i in range(min(count, 6)):
                    api_logger.set_phase(f"plans:clicking_plan_{i+1}")
                    btn = buttons.nth(i)
                    btn.scroll_into_view_if_needed(timeout=5_000)
                    reading_pause(1.0, 2.5)
                    bezier_click(page, btn)
                    reading_pause(1.5, 3.0)
                    close_any_modal(page)
                    print(f"    clicked plan {i+1}")
                return
        except Exception:
            continue

    print("  No plan buttons found to click — page may have been blocked.")
    reading_pause(2.0, 4.0)
    api_logger.set_phase("plans:scrolling_page")
    for _ in range(3):
        natural_scroll(page, "down", random.randint(200, 400))
        reading_pause(1.0, 2.0)


def close_any_modal(page: Page) -> None:
    for selector in [
        "button[aria-label*='Close' i]",
        "button:has-text('Close')",
        "button[aria-label$='modal Close']",
    ]:
        try:
            btn = page.locator(selector).first
            if btn.is_visible(timeout=2_000):
                bezier_click(page, btn, timeout=3_000)
                page.wait_for_timeout(500)
                return
        except Exception:
            continue
    try:
        page.keyboard.press("Escape")
        page.wait_for_timeout(500)
    except Exception:
        pass


def filter_excel(src_path: str, dst_path: str) -> int:
    wb = openpyxl.load_workbook(src_path)
    ws = wb["API Calls"]
    signal_col = None
    for col_idx in range(1, ws.max_column + 1):
        if ws.cell(row=1, column=col_idx).value == "Signal Level":
            signal_col = col_idx
            break
    if signal_col is None:
        signal_col = 9
    rows_to_delete = []
    for row_idx in range(ws.max_row, 1, -1):
        level = ws.cell(row=row_idx, column=signal_col).value
        if level in ("None", "Low"):
            rows_to_delete.append(row_idx)
    for row_idx in rows_to_delete:
        ws.delete_rows(row_idx, 1)
    wb.save(dst_path)
    return ws.max_row - 1


def run(playwright: Playwright, proxy: str | None = None) -> None:
    print(f"\n{'='*60}")
    print(f"  FRONTIER — LEVEL 4: Patchright (real Chrome TLS)")
    print(f"  headless=False | bezier_mouse=True | lognormal_typing=True")
    print(f"  warmup_browse=15s | proxy={'yes' if proxy else 'none'}")
    print(f"{'='*60}\n")

    launch_args: dict = {
        "headless": False,
        "channel": "chrome",
    }
    if proxy:
        launch_args["proxy"] = {"server": proxy}

    browser: Browser = playwright.chromium.launch(**launch_args)

    viewport = random.choice([
        {"width": 1920, "height": 1080},
        {"width": 1536, "height": 864},
        {"width": 1440, "height": 900},
    ])
    context_args: dict = {
        "viewport": viewport,
        "user_agent": (
            "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/128.0.0.0 Safari/537.36"
        ),
    }
    context: BrowserContext = browser.new_context(**context_args)
    page: Page = context.new_page()

    api_logger = ApiLogger()
    api_logger.attach(page)

    try:
        print(f"Opening {FRONTIER_URL}")
        api_logger.set_phase("landing_page:initial_load")
        page.goto(FRONTIER_URL, wait_until="domcontentloaded", timeout=45_000)
        reading_pause(2.0, 4.0)

        api_logger.snapshot_cookies(context, "01_after_landing_page_load")
        dismiss_cookie_banner(page)

        print("Warm-up browsing for ~15 seconds...")
        api_logger.set_phase("warmup:browsing")
        warmup_browse(page, duration_s=15.0)
        api_logger.snapshot_cookies(context, "02_after_warmup_browse")

        landing_ok = try_enter_address_landing(page, api_logger)

        if landing_ok:
            api_logger.snapshot_cookies(context, "03_after_landing_check_availability")
            api_logger.set_phase("redirect:waiting_for_buy_page")
            try:
                page.wait_for_url("**/buy**", timeout=15_000)
                print("  Redirected to buy page")
            except PatchrightTimeoutError:
                print("  No redirect to /buy — checking current page")
        else:
            print("  Landing page blocked — navigating to /buy directly")
            api_logger.set_phase("fallback:direct_buy_page")
            page.goto("https://frontier.com/buy", wait_until="domcontentloaded", timeout=45_000)

        reading_pause(2.0, 4.0)
        api_logger.snapshot_cookies(context, "04_after_buy_page_load")

        current_url = page.url
        if "/buy" in current_url:
            enter_address_buy_page(page, api_logger)
            api_logger.snapshot_cookies(context, "05_after_buy_page_check")

        api_logger.set_phase("plans_page:loading")
        reading_pause(3.0, 6.0)
        api_logger.snapshot_cookies(context, "06_after_plans_loaded")

        browse_plans(page, api_logger)
        api_logger.snapshot_cookies(context, "07_after_browsing_plans")

        print("\nDone. Closing browser.")
    except Exception as exc:
        print(f"\nERROR during automation: {exc}")
        api_logger.snapshot_cookies(context, "99_after_error")
    finally:
        time.sleep(1)
        os.makedirs(LOGS_DIR, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")

        full_path = os.path.join(LOGS_DIR, f"{LEVEL_NAME}_{ts}.xlsx")
        api_logger.save_to_excel(full_path)

        filtered_path = os.path.join(LOGS_DIR, f"{LEVEL_NAME}_filtered.xlsx")
        kept = filter_excel(full_path, filtered_path)
        print(f"Filtered to {kept} important calls -> {filtered_path}")

        context.close()
        browser.close()


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Level 4: Frontier automation with Patchright (real Chrome TLS)"
    )
    parser.add_argument("--proxy", type=str, default=None,
                        help="Proxy URL, e.g. http://user:pass@host:port")
    args = parser.parse_args()

    with sync_playwright() as p:
        run(p, proxy=args.proxy)


if __name__ == "__main__":
    main()
