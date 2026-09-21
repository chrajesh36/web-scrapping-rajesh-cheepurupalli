"""Level 4: Verizon broadband-plans automation using Patchright.

Patchright is a drop-in Playwright replacement that:
  - Drives the system-installed Chrome (matching real TLS/JA4 fingerprint)
  - Patches the CDP startup sequence to avoid automation markers
  - Keeps the same API as Playwright so all selectors work unchanged

Combined with:
  - behavior.py Bezier mouse / log-normal typing / natural scroll
  - Homepage warm-up browsing phase (15 s behavioral baseline)
  - Sensor POST interception and _abck state tracking

Usage:
    python verizon_patchright.py [--proxy URL]
"""

from __future__ import annotations

import argparse
import os
import random
import time
from datetime import datetime

import openpyxl
from patchright.sync_api import (
    Browser,
    BrowserContext,
    Page,
    Playwright,
    TimeoutError as PatchrightTimeoutError,
    sync_playwright,
)

from api_logger import ApiLogger
from behavior import (
    bezier_click,
    bezier_mouse_move,
    lognormal_type,
    natural_scroll,
    reading_pause,
    warmup_browse,
)

VERIZON_URL = "https://www.verizon.com/home/internet/"
ADDRESS = "140 West Street, New York, NY 10007"
LOGS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "logs")
LEVEL_NAME = "level4_patchright"


def dismiss_cookie_banner(page: Page) -> None:
    for selector in [
        "#truste-consent-button",
        "button:has-text('Accept All')",
        "button:has-text('Accept all')",
        "button.button:has-text('Close')",
    ]:
        try:
            for loc in page.locator(selector).all():
                if loc.is_visible():
                    loc.click(timeout=2_000)
                    page.wait_for_timeout(500)
                    return
        except Exception:
            continue


def enter_address_and_submit(page: Page, api_logger: ApiLogger) -> None:
    print("Entering service address...")
    api_logger.set_phase("address_entry:typing")

    address_field = page.locator(
        "#streetaddress, input[aria-label*='Street Address' i]"
    ).first
    address_field.wait_for(state="visible", timeout=15_000)
    bezier_click(page, address_field)
    address_field.fill("")

    lognormal_type(page, address_field, ADDRESS)
    reading_pause(1.0, 2.5)

    api_logger.set_phase("address_entry:autocomplete_suggestions")
    suggestion = page.locator(
        "[role='option'], ul[role='listbox'] li, div[class*='suggest' i] li"
    ).first
    try:
        suggestion.wait_for(state="visible", timeout=5_000)
        reading_pause(0.5, 1.5)
        bezier_click(page, suggestion)
        print("  selected suggestion")
    except PatchrightTimeoutError:
        print("  no autocomplete, pressing Enter")
        address_field.press("Enter")

    reading_pause(0.8, 2.0)

    api_logger.set_phase("address_entry:get_started_click")
    get_started = page.locator("button:has-text('Get started')").first
    get_started.wait_for(state="visible", timeout=15_000)
    bezier_click(page, get_started)
    print("  clicked 'Get started'")

    api_logger.set_phase("address_entry:business_residential_modal")
    residential_btn = page.locator("button:has-text('Continue as residential')").first
    try:
        residential_btn.wait_for(state="visible", timeout=6_000)
        reading_pause(0.5, 1.5)
        bezier_click(page, residential_btn)
        print("  selected 'Continue as residential'")
    except PatchrightTimeoutError:
        pass


def wait_for_offers_page(page: Page, api_logger: ApiLogger) -> None:
    print("Waiting for plans to load...")
    api_logger.set_phase("offers_page:loading")
    page.wait_for_url("**/inhome/buildproducts**", timeout=30_000)
    try:
        page.locator(
            "button.reviewDetailsbtn, button:has-text('Review details')"
        ).first.wait_for(state="visible", timeout=30_000)
    except PatchrightTimeoutError:
        pass
    reading_pause(2.0, 4.0)
    dismiss_cookie_banner(page)


def click_each_plan(page: Page, api_logger: ApiLogger) -> None:
    review_buttons = page.locator(
        "button.reviewDetailsbtn, button:has-text('Review details')"
    )
    total = review_buttons.count()
    if total == 0:
        print("  WARNING: No 'Review details' buttons found.")
        return
    print(f"Found {total} plan(s) on the offers page.")

    for index in range(total):
        print(f"\n  Opening plan {index + 1}/{total}...")
        api_logger.set_phase(f"plan_review:{index + 1}_of_{total}")
        try:
            btn = review_buttons.nth(index)
            btn.scroll_into_view_if_needed(timeout=15_000)
            reading_pause(1.0, 3.0)
            bezier_click(page, btn)
            reading_pause(1.5, 3.0)
            close_details_modal(page)
            print(f"  opened and closed plan {index + 1}")
        except Exception as exc:
            print(f"  could not open plan {index + 1}: {exc}")
            close_details_modal(page)
        reading_pause(1.0, 2.5)


def close_details_modal(page: Page) -> bool:
    for selector in [
        "button[aria-label$='modal Close']",
        "button[aria-label='Close']",
        "button:has-text('Close')",
    ]:
        try:
            btn = page.locator(selector).first
            if btn.is_visible(timeout=2_000):
                bezier_click(page, btn, timeout=3_000)
                page.wait_for_timeout(500)
                return True
        except Exception:
            continue
    try:
        page.keyboard.press("Escape")
        page.wait_for_timeout(500)
        return True
    except Exception:
        return False


def filter_excel(src_path: str, dst_path: str) -> int:
    wb = openpyxl.load_workbook(src_path)
    ws = wb["API Calls"]
    rows_to_delete = []
    for row_idx in range(ws.max_row, 1, -1):
        level = ws.cell(row=row_idx, column=9).value
        if level in ("None", "Low"):
            rows_to_delete.append(row_idx)
    for row_idx in rows_to_delete:
        ws.delete_rows(row_idx, 1)
    wb.save(dst_path)
    return ws.max_row - 1


def run(playwright: Playwright, proxy: str | None = None) -> None:
    print(f"\n{'='*60}")
    print(f"  LEVEL 4: Patchright (real Chrome TLS fingerprint)")
    print(f"  headless=False | bezier_mouse=True | lognormal_typing=True")
    print(f"  warmup_browse=15s | proxy={'yes' if proxy else 'none'}")
    print(f"{'='*60}\n")

    launch_args: dict = {
        "headless": False,
        "channel": "chrome",
    }
    if proxy:
        launch_args["proxy"] = {"server": proxy}

    browser: Browser = playwright.chromium.launch(**launch_args)

    viewport = random.choice([
        {"width": 1920, "height": 1080},
        {"width": 1536, "height": 864},
        {"width": 1440, "height": 900},
    ])
    context_args: dict = {
        "viewport": viewport,
        "user_agent": (
            "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/128.0.0.0 Safari/537.36"
        ),
    }
    context: BrowserContext = browser.new_context(**context_args)
    page: Page = context.new_page()

    api_logger = ApiLogger()
    api_logger.attach(page)

    try:
        print(f"Opening {VERIZON_URL}")
        api_logger.set_phase("landing_page:initial_load")
        page.goto(VERIZON_URL, wait_until="domcontentloaded", timeout=45_000)
        reading_pause(2.0, 4.0)

        api_logger.snapshot_cookies(context, "01_after_landing_page_load")
        dismiss_cookie_banner(page)

        print("Warm-up browsing for ~15 seconds...")
        api_logger.set_phase("warmup:browsing")
        warmup_browse(page, duration_s=15.0)
        api_logger.snapshot_cookies(context, "02_after_warmup_browse")

        enter_address_and_submit(page, api_logger)
        api_logger.snapshot_cookies(context, "03_after_get_started_click")

        wait_for_offers_page(page, api_logger)
        api_logger.snapshot_cookies(context, "04_after_offers_page_loaded")

        click_each_plan(page, api_logger)
        api_logger.snapshot_cookies(context, "05_after_all_plans_reviewed")
        print("\nDone. Closing browser.")
    except Exception as exc:
        print(f"\nERROR during automation: {exc}")
        api_logger.snapshot_cookies(context, "99_after_error")
    finally:
        time.sleep(1)
        os.makedirs(LOGS_DIR, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")

        full_path = os.path.join(LOGS_DIR, f"{LEVEL_NAME}_{ts}.xlsx")
        api_logger.save_to_excel(full_path)

        filtered_path = os.path.join(LOGS_DIR, f"{LEVEL_NAME}_filtered.xlsx")
        kept = filter_excel(full_path, filtered_path)
        print(f"Filtered to {kept} important calls -> {filtered_path}")

        context.close()
        browser.close()


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Level 4: Verizon automation with Patchright (real Chrome TLS)"
    )
    parser.add_argument("--proxy", type=str, default=None,
                        help="Proxy URL, e.g. http://user:pass@host:port")
    args = parser.parse_args()

    with sync_playwright() as p:
        run(p, proxy=args.proxy)


if __name__ == "__main__":
    main()
